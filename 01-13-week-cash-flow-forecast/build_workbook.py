from pathlib import Path
import csv

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "outputs" / "forecast_results.csv"
RAW = ROOT / "data" / "raw"
WORKBOOK = ROOT / "excel" / "13_week_cash_forecast.xlsx"


BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
RED = "F4CCCC"
GREEN = "D9EAD3"


def read_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as file:
        return list(csv.DictReader(file))


def write_table(sheet, rows, table_name):
    if not rows:
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row[column] for column in headers])
    last_row = sheet.max_row
    last_col = sheet.max_column
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"


def style_sheet(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center")
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 28)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def main():
    rows = read_csv(OUTPUT)
    base = [row for row in rows if row["scenario"] == "Base"]
    stress = [row for row in rows if row["scenario"] != "Base"]
    workbook = Workbook()
    control = workbook.active
    control.title = "Control Panel"
    control.append(["13-Week Cash Flow Forecast", "Value"])
    for key, value in [("Company", "Pragati Distribution Pvt Ltd"), ("Currency", "INR"), ("Forecast Horizon", 13), ("Minimum Cash Threshold", 1000000), ("Stress Scenario", "Customers pay 20 days late")]:
        control.append([key, value])
    control["A8"] = "Refresh instructions"
    control["A8"].font = Font(bold=True, color="FFFFFF")
    control["A8"].fill = PatternFill("solid", fgColor=BLUE)
    control["A9"] = "Replace files in data/raw, refresh Power Query, then refresh this workbook."
    control.merge_cells("A9:B9")

    for name, dataset, table_name in [("Base Forecast", base, "BaseForecast"), ("Stress Forecast", stress, "StressForecast")]:
        sheet = workbook.create_sheet(name)
        write_table(sheet, dataset, table_name)
        style_sheet(sheet)
    raw_sources = [
        ("Raw Invoices", "customer_invoices.csv", "RawInvoices"),
        ("Raw Supplier Payments", "supplier_payments.csv", "RawSupplierPayments"),
        ("Raw Other Cash Flows", "other_cash_flows.csv", "RawOtherCashFlows"),
        ("Raw Configuration", "forecast_config.csv", "RawConfiguration"),
        ("Raw Opening Cash", "opening_cash.csv", "RawOpeningCash"),
    ]
    for name, filename, table_name in raw_sources:
        sheet = workbook.create_sheet(name)
        write_table(sheet, read_csv(RAW / filename), table_name)
        style_sheet(sheet)
    comparison = workbook.create_sheet("Scenario Comparison")
    comparison.append(["Week Start", "Base Closing Cash", "Stress Closing Cash", "Stress Cash Gap", "Base Alert", "Stress Alert"])
    for index, row in enumerate(base, start=2):
        comparison.append([row["week_start"], f"='Base Forecast'!F{index}", f"='Stress Forecast'!F{index}", f"=B{index}-C{index}", f"='Base Forecast'!G{index}", f"='Stress Forecast'!G{index}"])
    style_sheet(comparison)
    chart = LineChart()
    chart.title = "Base versus Stress Closing Cash"
    chart.y_axis.title = "INR"
    chart.x_axis.title = "Week"
    chart.add_data(Reference(comparison, min_col=2, max_col=3, min_row=1, max_row=comparison.max_row), titles_from_data=True)
    chart.set_categories(Reference(comparison, min_col=1, min_row=2, max_row=comparison.max_row))
    comparison.add_chart(chart, "H2")

    receipts = workbook.create_sheet("Receipts")
    write_table(receipts, [{"week_start": row["week_start"], "customer_receipts": row["receipts"]} for row in base], "Receipts")
    style_sheet(receipts)
    payments = workbook.create_sheet("Payments")
    write_table(payments, [{"week_start": row["week_start"], "supplier_and_other_payments": row["payments"]} for row in base], "Payments")
    style_sheet(payments)
    quality = workbook.create_sheet("Data Quality")
    quality.append(["Check", "Result"])
    quality.append(["Base forecast weeks", len(base)])
    quality.append(["Stress forecast weeks", len(stress)])
    quality.append(["Scenario week counts match", "PASS" if len(base) == len(stress) == 13 else "FAIL"])
    quality.append(["Base closing cash values present", "PASS" if all(row["closing_cash"] for row in base) else "FAIL"])
    quality.append(["Stress closing cash values present", "PASS" if all(row["closing_cash"] for row in stress) else "FAIL"])
    style_sheet(quality)
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    cell.number_format = '#,##0.00'
    control["B5"].number_format = '#,##0'
    control["B4"].number_format = '0'
    WORKBOOK.parent.mkdir(exist_ok=True)
    try:
        workbook.save(WORKBOOK)
        saved_path = WORKBOOK
    except PermissionError:
        saved_path = WORKBOOK.with_name("13_week_cash_forecast_with_raw_data.xlsx")
        workbook.save(saved_path)
    print(f"Created {saved_path}")


if __name__ == "__main__":
    main()
